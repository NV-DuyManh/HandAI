from pathlib import Path

from docx import Document


OUT = Path(__file__).parent
MAIN = OUT / "HandAI_Final_Live_Data.docx"
CHANGE = OUT / "Bao_cao_cap_nhat_du_lieu_thuc_HandAI.docx"


def replace_start(document: Document, prefix: str, replacement: str) -> None:
    paragraph = next((p for p in document.paragraphs if p.text.startswith(prefix)), None)
    if paragraph is None:
        raise RuntimeError(f"Missing paragraph prefix: {prefix}")
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)


main = Document(MAIN)
replace_start(
    main,
    "The research screen no longer supplies",
    "The research screen no longer supplies headline accuracy, dataset totals, comparison gains, latency or confidence through "
    "fallback constants. When no verified live lines exist, metric cards show an em dash or an explanatory empty state. When "
    "verified lines exist, every displayed value is derived from the filtered session and line records. The recognition-detail "
    "view also renders only the input-image URI saved with that session; if the file is unavailable, it shows an explicit message "
    "instead of generating a notebook preview. [4, 10]",
)
replace_start(
    main,
    "[10] apps/mobile/src/app/handai-analytics.tsx",
    "[10] apps/mobile/src/app/handai-analytics.tsx, evaluation-history.tsx, (tabs)/index.tsx, crop.tsx, "
    "ocr-pilot/multiline-review.tsx and multiline-result.tsx; screens/HandAiTrialAnalyticsScreen.tsx; "
    "services/api/ocrPilotService.ts; services/analytics/handAiDashboardMetrics.ts; components/ocr/OCRProgressLoader.tsx. "
    "Current routing, line editing, stored-image evidence, polling, progress and live analytics states.",
)
main.save(MAIN)


change = Document(CHANGE)
replace_start(
    change,
    "Các phiên mẫu dựng sẵn",
    "Các phiên mẫu dựng sẵn không còn xuất hiện trong dashboard nghiên cứu và lịch sử. Trang chi tiết chỉ mở phiên COMPLETED "
    "không phải dữ liệu mẫu; mọi số liệu, chuỗi OCR, gợi ý, kết quả cuối và lỗi đều lấy từ lineMetrics/errorRecords của đúng phiên. "
    "Ảnh được lấy từ URI camera/gallery đã lưu khi hoàn tất; nếu tệp không tồn tại, giao diện báo không có ảnh và không vẽ ảnh vở "
    "thay thế. Khai báo Stack.Screen ocr-pilot không có route index cũng đã được xóa để chấm dứt cảnh báo Layout children.",
)
replace_start(
    change,
    "Giữ CER 11,3388%",
    "Giữ CER 11,3388% trong model manifest như bằng chứng bàn giao mô hình; không trộn số này với dashboard phiên trực tiếp. "
    "Báo cáo đồng thời ghi rõ trang chi tiết chỉ hiển thị ảnh đầu vào đã lưu, dữ liệu của đúng phiên và trạng thái chưa có dữ liệu "
    "khi thiếu bằng chứng.",
)
replace_start(
    change,
    "TypeScript hoàn tất không lỗi.",
    "TypeScript hoàn tất không lỗi. Bốn bộ kiểm thử HandAI liên quan có 100/100 ca đạt. Các ca mới kiểm tra việc loại dữ liệu mẫu, "
    "tính số từ dòng xác minh, không dựng phần trăm khi thiếu reference, và hiển thị đúng URI ảnh cùng nội dung của phiên trực tiếp. "
    "Không huấn luyện mô hình và không commit hoặc push trong lần cập nhật này.",
)
replace_start(
    change,
    "documents\n",
    "documents\n"
    "SKILL.md: C:/Users/Admin/.codex/plugins/cache/openai-primary-runtime/documents/26.929.10730/skills/documents/SKILL.md\n"
    "Why selected: yêu cầu render và kiểm tra trực quan trước khi bàn giao.\n"
    "Applied to: xuất DOCX, render từng trang và kiểm tra bố cục.",
)
change.save(CHANGE)

print(f"Updated {MAIN}")
print(f"Updated {CHANGE}")
