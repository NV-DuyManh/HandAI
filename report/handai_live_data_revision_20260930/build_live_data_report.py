from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, "E:/HandAI/report/handai_visual_revision_20260930/python_deps")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("E:/HandAI")
OUT = Path(__file__).parent
FIGURES = OUT / "figures"
FIGURES.mkdir(parents=True, exist_ok=True)
SOURCE_DOC = ROOT / "report/handai_current_app_revision_20260930/HandAI_Final.docx"
PROCESSING_SCREEN = Path(
    "C:/Users/Admin/Downloads/1790754161514_121714083976711462_6661291260689322471_4aa1b437d2d928a64e8ac7443a03e6d0.jpg"
)

NAVY = "#203D61"
BLUE = "#2563EB"
TEAL = "#0F8A72"
PURPLE = "#7C3AED"
GRAY = "#64748B"
PALE = "#F5F8FC"

plt.rcParams.update(
    {
        "font.family": "DejaVu Sans",
        "font.size": 10,
        "text.color": NAVY,
        "savefig.facecolor": "white",
    }
)


def rounded_box(ax, x, y, w, h, title, body, color=BLUE, title_size=10, body_size=8.5):
    ax.add_patch(
        FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.025,rounding_size=0.06",
            facecolor=PALE,
            edgecolor=color,
            linewidth=1.7,
        )
    )
    ax.text(x + w / 2, y + h * 0.69, title, ha="center", va="center", weight="bold", fontsize=title_size, color=color)
    ax.text(x + w / 2, y + h * 0.34, body, ha="center", va="center", fontsize=body_size, linespacing=1.35, color=GRAY)


def arrow(ax, start, end, color=GRAY):
    ax.annotate("", xy=end, xytext=start, arrowprops={"arrowstyle": "-|>", "lw": 1.6, "color": color})


def make_live_pipeline(path: Path):
    fig, ax = plt.subplots(figsize=(4.4, 8.0))
    ax.set(xlim=(0, 10), ylim=(0, 16))
    ax.axis("off")
    ax.text(5, 15.25, "LIVE ANALYTICS DATA PATH", ha="center", weight="bold", fontsize=13, color=NAVY)
    ax.text(5, 14.72, "No fallback benchmark values", ha="center", fontsize=9.5, color=GRAY)
    boxes = [
        (12.5, "COMPLETED SESSIONS", "Recognition history stored\non the current device", BLUE),
        (9.7, "SESSION FILTER", "Keep status COMPLETED\nExclude isSampleData = true", TEAL),
        (6.9, "VERIFIED LINES", "Use EXPLICIT or USER_CONFIRMED\nnon-empty reference text", PURPLE),
        (4.1, "CORPUS AGGREGATION", "Sum edit counts and reference units\nCount exact raw and final matches", BLUE),
        (1.3, "DASHBOARD OUTPUT", "Show computed value\nor an explicit empty state", TEAL),
    ]
    for y, title, body, color in boxes:
        rounded_box(ax, 1.05, y, 7.9, 1.8, title, body, color)
    for y1, y2 in [(12.5, 11.5), (9.7, 8.7), (6.9, 5.9), (4.1, 3.1)]:
        arrow(ax, (5, y1), (5, y2))
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)


def make_metric_scope(path: Path):
    fig, ax = plt.subplots(figsize=(4.5, 7.2))
    ax.set(xlim=(0, 10), ylim=(0, 14))
    ax.axis("off")
    ax.text(5, 13.35, "METRIC INCLUSION RULES", ha="center", weight="bold", fontsize=13, color=NAVY)
    rounded_box(ax, 0.8, 10.6, 8.4, 1.8, "ALL LIVE COMPLETED SESSIONS", "Session count, processed lines\nand mean session latency", BLUE)
    arrow(ax, (5, 10.55), (5, 9.75))
    rounded_box(ax, 0.8, 7.7, 8.4, 1.8, "VERIFIED REFERENCE LINES", "Explicit or user-confirmed reference\nSample sessions remain excluded", PURPLE)
    arrow(ax, (5, 7.65), (5, 6.85))
    rounded_box(ax, 0.8, 4.3, 4.0, 1.9, "RAW OCR", "Exact-line accuracy\nCorpus CER and WER", TEAL)
    rounded_box(ax, 5.2, 4.3, 4.0, 1.9, "FINAL SELECTION", "Exact-line accuracy\nDecision-source counts", TEAL)
    arrow(ax, (2.8, 4.25), (2.8, 3.5))
    arrow(ax, (7.2, 4.25), (7.2, 3.5))
    rounded_box(ax, 0.8, 1.0, 8.4, 1.85, "PAIRED WORKFLOW CHANGE", "Final line accuracy minus raw line accuracy\nreported in percentage points", BLUE)
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)


def make_display_logic(path: Path):
    fig, ax = plt.subplots(figsize=(4.5, 7.2))
    ax.set(xlim=(0, 10), ylim=(0, 14))
    ax.axis("off")
    ax.text(5, 13.35, "DASHBOARD DISPLAY LOGIC", ha="center", weight="bold", fontsize=13, color=NAVY)
    rounded_box(ax, 0.8, 10.8, 8.4, 1.65, "LOAD LOCAL HISTORY", "Refresh reads the analytics store again", BLUE)
    arrow(ax, (5, 10.75), (5, 9.85))
    rounded_box(ax, 1.6, 7.9, 6.8, 1.5, "ANY VERIFIED LINES?", "Same inclusion rules for every metric", PURPLE)
    ax.text(2.0, 7.35, "NO", ha="center", weight="bold", color=GRAY)
    ax.text(8.0, 7.35, "YES", ha="center", weight="bold", color=TEAL)
    arrow(ax, (3.0, 7.9), (2.1, 6.65))
    arrow(ax, (7.0, 7.9), (7.9, 6.65))
    rounded_box(ax, 0.35, 4.25, 4.1, 1.9, "EMPTY STATE", "Show — and explain\nhow verified data\nis created", GRAY)
    rounded_box(ax, 5.55, 4.25, 4.1, 1.9, "COMPUTED STATE", "Show metric value,\ncoverage and\nprovenance", TEAL)
    arrow(ax, (2.4, 4.2), (3.8, 3.25))
    arrow(ax, (7.6, 4.2), (6.2, 3.25))
    rounded_box(ax, 1.1, 1.0, 7.8, 1.65, "NO SEEDED NUMBERS", "No 88.2%, 80 to 92, 500-line\nor latency fallback is rendered", BLUE)
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)


def make_architecture(path: Path):
    fig, ax = plt.subplots(figsize=(9.2, 5.2))
    ax.set(xlim=(0, 10), ylim=(0, 5.5))
    ax.axis("off")
    rounded_box(ax, 0.1, 3.75, 2.8, 1.15, "MOBILE CLIENT", "Capture, crop, review\nraw, suggested and final text", BLUE)
    rounded_box(ax, 3.6, 3.75, 2.8, 1.15, "SPRING BACKEND", "Trials, crops, feedback\nand advisor orchestration", BLUE)
    rounded_box(ax, 7.1, 3.75, 2.8, 1.15, "PYTHON AI SERVICE", "OpenCV plus CRNN and CTC\ncanonical fixture branch", BLUE)
    arrow(ax, (2.95, 4.32), (3.55, 4.32), BLUE)
    arrow(ax, (6.45, 4.32), (7.05, 4.32), BLUE)
    rounded_box(ax, 0.1, 1.6, 2.8, 1.15, "LIVE ANALYTICS", "Filter completed non-sample sessions\naggregate verified lines", TEAL)
    rounded_box(ax, 3.6, 1.6, 2.8, 1.15, "PERSISTENCE", "PostgreSQL trials and lines\nMinIO pages and crops", TEAL)
    rounded_box(ax, 7.1, 1.6, 2.8, 1.15, "OPTIONAL ADVISORS", "Groq or Gemini\nLocal-Advisor fallback", PURPLE)
    arrow(ax, (1.5, 3.7), (1.5, 2.8), TEAL)
    arrow(ax, (5.0, 3.7), (5.0, 2.8), BLUE)
    arrow(ax, (8.5, 3.7), (8.5, 2.8), PURPLE)
    ax.text(5, 0.8, "Dashboard values come from the device history store; missing verified evidence remains visibly unavailable.", ha="center", fontsize=10, color=GRAY)
    ax.text(5, 0.35, "CRNN inference remains server-side in the inspected implementation.", ha="center", weight="bold", fontsize=10.5, color=NAVY)
    fig.savefig(path, dpi=240, bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)


LIVE_PIPELINE = FIGURES / "live_analytics_pipeline.png"
METRIC_SCOPE = FIGURES / "metric_inclusion_rules.png"
DISPLAY_LOGIC = FIGURES / "dashboard_display_logic.png"
ARCHITECTURE = FIGURES / "architecture_live_metrics.png"
make_live_pipeline(LIVE_PIPELINE)
make_metric_scope(METRIC_SCOPE)
make_display_logic(DISPLAY_LOGIC)
make_architecture(ARCHITECTURE)


doc = Document(SOURCE_DOC)
changes: list[dict[str, str]] = []


def set_paragraph_text(paragraph, text: str):
    old = paragraph.text
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    changes.append({"before": old, "after": text})


def replace_start(prefix: str, text: str):
    paragraph = next(p for p in doc.paragraphs if p.text.startswith(prefix))
    set_paragraph_text(paragraph, text)
    return paragraph


replacements = {
    "2.1.1 Processing and dashboard interface": "2.1.1 Processing and live analytics interface",
    "The processing screen exposes": (
        "The processing screen exposes five stages: image preparation, line detection, OCR, AI verification and completion. "
        "Its 35% indicator represents workflow progress rather than recognition quality. The live analytics dashboard now reads "
        "completed device sessions, excludes bundled sample records and calculates values only where usable reference text exists. [4, 10]"
    ),
    "Figure 4.": (
        "Figure 4. User-provided processing screen and code-derived live analytics data path. The progress indicator describes "
        "workflow state; the dashboard path shows how runtime evidence replaces fixed presentation values. [4, 7, 10]"
    ),
    "HandAI separates the mobile client": (
        "HandAI separates the mobile client, Java backend, Python AI service, persistence and client analytics. CRNN inference runs "
        "in the AI service. PostgreSQL holds trials and line records; MinIO holds page images and line crops. The research dashboard "
        "loads the client history store, filters completed non-sample sessions and aggregates verified line records. [2–5, 10]"
    ),
    "Figure 7.": (
        "Figure 7. Current implementation architecture with background AI advice and a live dashboard aggregation layer. Sources: "
        "model code, OCR API, multiline service and mobile analytics. [1–5, 10]"
    ),
    "The store currently averages available session CER/WER values": (
        "The dedicated dashboard aggregator filters completed non-sample sessions before display. It computes corpus CER and WER "
        "from total edit counts divided by total reference units, counts exact raw and final line matches, and derives confidence, "
        "latency and error items from the same live records. The legacy global analytics method remains separate and is not used for "
        "the research dashboard. [4, 10]"
    ),
    "The primary recorded model result is": (
        "The primary recorded model result remains validation CER 11.3388% at checkpoint step 16,900 on the declared 500-sample "
        "validation split. This model-handoff measurement remains separate from device-session analytics, whose values vary with the "
        "completed live records stored on the device. [1, 4, 10]"
    ),
    "The current research screen hard-codes": (
        "The research screen no longer supplies headline accuracy, dataset totals, comparison gains, latency or confidence through "
        "fallback constants. When no verified live lines exist, metric cards show an em dash or an explanatory empty state. When "
        "verified lines exist, every displayed value is derived from the filtered session and line records. The recognition-detail "
        "view also renders only the input-image URI saved with that session; if the file is unavailable, it shows an explicit message "
        "instead of generating a notebook preview. [4, 10]"
    ),
    "The 80% → 92% comparison": (
        "The dashboard reports the change from raw exact-line accuracy to the selected final result in percentage points. Because the "
        "final selection can reflect both an AI suggestion and manual review, the interface labels this value as assisted workflow "
        "improvement rather than attributing the entire change to AI alone. [4, 10]"
    ),
    "Current research dashboard views": "Live analytics calculation and display states",
    "Figure 10.": (
        "Figure 10. Metric inclusion rules and dashboard display logic derived from the current implementation. Verified values use "
        "completed non-sample sessions; unavailable evidence produces an explicit empty state. [4, 10]"
    ),
    "These panels communicate the intended progression": (
        "These diagrams document the operational dashboard calculation. A user-confirmed reference is suitable for workflow monitoring "
        "but may not be independent of the selected output; the matched research comparison in Section 6.6 still requires frozen "
        "external references and exported line-level results."
    ),
    "The client maintains dataset/model registrations": (
        "The client maintains dataset/model registrations, trial analytics, history and exports; the backend persists multiline trials, "
        "line crops and feedback. Completed sessions enter the local history, capped at 50 sessions. Native storage uses SecureStore "
        "and web storage uses localStorage, with an in-memory fallback. The research dashboard excludes sessions marked as sample data "
        "and displays only values derived from the remaining records. [4, 5, 10]"
    ),
    "[4] apps/mobile/": (
        "[4] apps/mobile/src/services/analytics/handAiAnalyticsStore.ts and handAiDashboardMetrics.ts. Trial scoring, reference-status "
        "rules, persistence, live-session filtering and corpus metric aggregation."
    ),
    "[7] User-supplied app screenshots": (
        "[7] User-supplied app screenshots: capture/crop/review/result screens from the source document and the current processing "
        "screen. These document UI behavior; the analytics diagrams were generated from the inspected source code."
    ),
    "[10] apps/mobile/src/app/handai-analytics.tsx": (
        "[10] apps/mobile/src/app/handai-analytics.tsx, evaluation-history.tsx, (tabs)/index.tsx, crop.tsx, "
        "ocr-pilot/multiline-review.tsx and multiline-result.tsx; screens/HandAiTrialAnalyticsScreen.tsx; "
        "services/api/ocrPilotService.ts; services/analytics/handAiDashboardMetrics.ts; components/ocr/OCRProgressLoader.tsx. "
        "Current routing, line editing, stored-image evidence, polling, progress and live analytics states."
    ),
    "Score A, B and C against the same reference.": (
        "Score A, B and C against the same reference. Compute corpus CER/WER from total edit counts and total reference units; report "
        "exact-line matches with numerator and denominator. Report newly corrected and newly damaged lines, provider/fallback counts "
        "and paired uncertainty intervals. Express the raw-to-final difference in percentage points and report relative change only "
        "when its denominator and interpretation are explicit."
    ),
}

for prefix, value in replacements.items():
    replace_start(prefix, value)


scope_table = next(table for table in doc.tables if table.cell(0, 0).text == "Scope")
scope_table.cell(2, 1).text = "Legacy cross-session summary; retained for existing exports and trial views."
scope_table.cell(3, 1).text = "Filter completed non-sample sessions and compute verified line metrics from the local history store."

evidence_table = next(table for table in doc.tables if table.cell(0, 0).text == "Evidence")
while len(evidence_table.rows) > 1:
    evidence_table._tbl.remove(evidence_table.rows[-1]._tr)
evidence_rows = [
    ("Checkpoint manifest", "Validation CER 11.3388%", "Existing model-handoff evidence"),
    ("Live completed sessions", "Non-sample records on the current device", "Session count, processed lines, latency and history"),
    ("Verified line records", "Explicit or user-confirmed reference", "Raw/final line accuracy, corpus CER/WER, confidence and errors"),
    ("No usable reference", "Value unavailable", "Em dash or explanatory empty state; no fallback score"),
]
for row_values in evidence_rows:
    cells = evidence_table.add_row().cells
    for cell, value in zip(cells, row_values):
        cell.text = value


def drawing_rids(table):
    result = []
    for row in table.rows:
        for cell in row.cells:
            result.extend(cell._tc.xpath(".//a:blip/@r:embed"))
    return result


def clear_cell(cell):
    cell._tc.clear_content()


def set_image_cell(cell, image_path: Path, label: str, height: float):
    clear_cell(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), height=Inches(height))
    label_p = cell.add_paragraph(label)
    label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label_p.paragraph_format.keep_with_next = True
    for run in label_p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)


figure4_table = next(table for table in doc.tables if "(a) Processing status" in " ".join(c.text for c in table.rows[0].cells))
figure10_table = next(table for table in doc.tables if "(a) AI assistance presentation" in " ".join(c.text for c in table.rows[0].cells))
old_rids = drawing_rids(figure4_table) + drawing_rids(figure10_table)

set_image_cell(figure4_table.cell(0, 0), PROCESSING_SCREEN, "(a) Processing status", 5.35)
set_image_cell(figure4_table.cell(0, 1), LIVE_PIPELINE, "(b) Live metric data path", 5.35)
set_image_cell(figure10_table.cell(0, 0), METRIC_SCOPE, "(a) Metric inclusion rules", 4.65)
set_image_cell(figure10_table.cell(0, 1), DISPLAY_LOGIC, "(b) Dashboard display logic", 4.65)

for rid in old_rids:
    if not doc.element.body.xpath(f'.//a:blip[@r:embed="{rid}"]'):
        doc.part.drop_rel(rid)


old_architecture = ROOT / "report/handai_current_app_revision_20260930/figures/architecture_current.png"
if old_architecture.exists():
    old_hash = hashlib.sha256(old_architecture.read_bytes()).hexdigest()
    for part in doc.part.related_parts.values():
        if part.content_type.startswith("image/") and hashlib.sha256(part.blob).hexdigest() == old_hash:
            part._blob = ARCHITECTURE.read_bytes()


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9"):
    props = cell._tc.get_or_add_tcPr()
    borders = props.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), color)


for table in (scope_table, evidence_table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(table.rows):
        tr_props = row._tr.get_or_add_trPr()
        if tr_props.find(qn("w:cantSplit")) is None:
            tr_props.append(OxmlElement("w:cantSplit"))
        if row_index == 0 and tr_props.find(qn("w:tblHeader")) is None:
            tr_props.append(OxmlElement("w:tblHeader"))
        for cell in row.cells:
            set_cell_borders(cell)
            set_cell_shading(cell, "365F91" if row_index == 0 else ("F4F7FB" if row_index % 2 == 0 else "FFFFFF"))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10.5)
                    run.bold = row_index == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else "1E293B")

for paragraph in doc.paragraphs:
    if re.match(r"^Figure \d+\.", paragraph.text):
        paragraph.style = doc.styles["Figure Caption"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_together = True
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            run.italic = True

for doc_pr in doc.element.body.xpath(".//wp:docPr"):
    paragraph_node = doc_pr
    while paragraph_node is not None and paragraph_node.tag != qn("w:p"):
        paragraph_node = paragraph_node.getparent()
    if paragraph_node is None:
        continue
    cell_node = paragraph_node.getparent()
    candidates = []
    if cell_node.tag == qn("w:tc"):
        candidates = ["".join(n.text or "" for n in p.iter(qn("w:t"))) for p in cell_node.iter(qn("w:p"))]
    labels = [text for text in candidates if text.startswith("(a)") or text.startswith("(b)")]
    if labels:
        doc_pr.set("descr", labels[0])

settings = doc.settings.element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

MAIN_OUTPUT = OUT / "HandAI_Final_Live_Data.docx"
doc.save(MAIN_OUTPUT)


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


report = Document()
section = report.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
normal = report.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for style_name in ("Title", "Heading 1", "Heading 2"):
    style = report.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.color.rgb = RGBColor(0, 0, 0)

report.add_paragraph("Báo cáo cập nhật dữ liệu thật cho HandAI", style="Title")
report.add_paragraph(
    "Bản cập nhật ngày 30/09/2026 sửa cả mã dashboard và báo cáo Word. Dashboard nghiên cứu hiện lấy dữ liệu từ lịch sử nhận dạng "
    "thực trên thiết bị, loại bản ghi mẫu và không hiển thị các phần trăm dự phòng khi chưa có dòng được xác minh."
)

report.add_heading("Thay đổi trong mã ứng dụng", level=1)
for item in (
    "Thêm handAiDashboardMetrics.ts để tạo một nguồn tính duy nhất cho dashboard từ các phiên COMPLETED và isSampleData khác true.",
    "Dòng được tính accuracy, CER và WER phải có evaluationStatus EVALUATED, groundTruthStatus EXPLICIT hoặc USER_CONFIRMED, và ground truth không rỗng.",
    "CER và WER được tính theo tổng edit distance chia tổng ký tự hoặc token tham chiếu trên toàn bộ dòng hợp lệ; không lấy trung bình đơn giản giữa các phiên.",
    "Các thẻ accuracy, confidence, latency, lịch sử và phân tích lỗi dùng bản ghi thật. Khi chưa đủ dữ liệu, giao diện hiển thị dấu gạch ngang hoặc trạng thái chờ dữ liệu xác minh.",
    "Chênh lệch raw-to-final được ghi là Assisted Workflow Improvement theo điểm phần trăm vì kết quả cuối có thể gồm cả gợi ý AI và chỉnh tay.",
    "Các phiên mẫu dựng sẵn không còn xuất hiện trong dashboard nghiên cứu và lịch sử của màn hình này.",
    "Trang chi tiết nhận dạng chỉ mở phiên COMPLETED không phải dữ liệu mẫu; mọi số liệu, chuỗi OCR, gợi ý, kết quả cuối và lỗi đều lấy từ lineMetrics/errorRecords của đúng phiên.",
    "Ảnh trong trang chi tiết lấy từ URI camera/gallery đã chuyển sang RecognitionSession khi hoàn tất. Nếu URI không tồn tại hoặc tệp đã mất, giao diện báo không có ảnh và không vẽ ảnh vở thay thế.",
    "Xóa khai báo Stack.Screen ocr-pilot không có route index, chấm dứt cảnh báo [Layout children] trên thiết bị.",
):
    add_bullet(report, item)

report.add_heading("Nguồn và phạm vi số liệu", level=1)
table = report.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ("Phạm vi", "Dữ liệu sử dụng", "Khi thiếu dữ liệu")
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text
rows = (
    ("Phiên", "Phiên hoàn thành, không phải dữ liệu mẫu", "Hiển thị 0 phiên"),
    ("Accuracy dòng", "So khớp chính xác OCR thô và kết quả cuối với reference", "Hiển thị —"),
    ("CER và WER", "Tổng edit distance trên tổng đơn vị reference", "Hiển thị —"),
    ("Confidence và lỗi", "Các dòng đã xác minh trong phiên thật", "Hiển thị trạng thái chờ"),
    ("Latency", "processingTimeSeconds của phiên hoàn thành", "Hiển thị —"),
)
for values in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, values):
        cell.text = text
for row_index, row in enumerate(table.rows):
    for cell in row.cells:
        set_cell_borders(cell)
        set_cell_shading(cell, "365F91" if row_index == 0 else ("F4F7FB" if row_index % 2 == 0 else "FFFFFF"))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)
                run.bold = row_index == 0
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else "1E293B")

report.add_heading("Thay đổi trong báo cáo Word", level=1)
for item in (
    "Bỏ ba ảnh dashboard chứa số cố định cũ và bỏ toàn bộ diễn giải coi các số đó là giao diện hiện hành.",
    "Giữ ảnh tiến trình xử lý do chủ dự án cung cấp vì ảnh này vẫn mô tả đúng năm bước của app.",
    "Thay phần dashboard bằng sơ đồ luồng dữ liệu, quy tắc chọn dòng và logic hiển thị được dựng trực tiếp từ mã mới.",
    "Cập nhật kiến trúc, mục 6.1, 6.3, mục theo dõi thí nghiệm và danh mục nguồn để phản ánh bộ tổng hợp dashboard mới.",
    "Giữ CER 11,3388% trong model manifest như bằng chứng bàn giao mô hình; không trộn số này với dashboard phiên trực tiếp.",
    "Bổ sung mô tả trang chi tiết nhận dạng mới: chỉ hiển thị ảnh đầu vào đã lưu, dữ liệu của đúng phiên và trạng thái chưa có dữ liệu khi thiếu bằng chứng.",
):
    add_bullet(report, item)

report.add_heading("Kiểm chứng", level=1)
report.add_paragraph(
    "TypeScript hoàn tất không lỗi. Bốn bộ kiểm thử HandAI liên quan có 100/100 ca đạt. Các ca mới kiểm tra việc loại dữ liệu mẫu, "
    "tính số từ dòng xác minh, không dựng phần trăm khi thiếu reference, và hiển thị đúng URI ảnh cùng nội dung của phiên trực tiếp. "
    "Không huấn luyện mô hình và không commit hoặc push trong lần cập nhật này."
)
report.add_paragraph(
    "USER_CONFIRMED hỗ trợ giám sát vận hành nhưng không mặc nhiên là ground truth độc lập, vì người dùng có thể xác nhận chính kết quả đang xem. "
    "Một benchmark nghiên cứu vẫn cần reference đóng băng từ trước và xuất kết quả theo từng dòng để so sánh raw OCR, AI-only và kết quả cuối trên cùng mẫu."
)

report.add_heading("Skills Applied", level=1)
report.add_paragraph(
    "vercel-react-best-practices\n"
    "SKILL.md: .agents/skills/react-best-practices/SKILL.md\n"
    "Why selected: cập nhật React Native dashboard và luồng tải dữ liệu.\n"
    "Applied to: memoization, tải phiên trực tiếp, trạng thái rỗng, giữ URI ảnh và kiểm thử màn hình."
)
report.add_paragraph(
    "docx\n"
    "SKILL.md: .agents/skills/docx/SKILL.md\n"
    "Why selected: chỉnh tài liệu Word hiện hữu mà vẫn giữ cấu trúc học thuật.\n"
    "Applied to: nội dung, bảng, hình và chú thích."
)
report.add_paragraph(
    "documents\n"
    "SKILL.md: C:/Users/Admin/.codex/plugins/cache/openai-primary-runtime/documents/26.929.10730/skills/documents/SKILL.md\n"
    "Why selected: yêu cầu render và kiểm tra trực quan trước khi bàn giao.\n"
    "Applied to: xuất DOCX, render từng trang và kiểm tra bố cục."
)

for element in (report.element, report.styles.element):
    for border in list(element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)

CHANGE_REPORT = OUT / "Bao_cao_cap_nhat_du_lieu_thuc_HandAI.docx"
report.save(CHANGE_REPORT)

manifest = {
    "source": str(SOURCE_DOC),
    "outputs": [str(MAIN_OUTPUT), str(CHANGE_REPORT)],
    "changes": changes,
    "removed_dashboard_image_relationships": old_rids,
    "figures": [str(LIVE_PIPELINE), str(METRIC_SCOPE), str(DISPLAY_LOGIC), str(ARCHITECTURE)],
}
(OUT / "changes.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Saved {MAIN_OUTPUT}")
print(f"Saved {CHANGE_REPORT}")
